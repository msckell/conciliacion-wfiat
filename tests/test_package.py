"""Package build and Excel, offline, with a tiny hand built engine output."""

from openpyxl import load_workbook

from cierre.config import Chain, Deployment, Token
from cierre.package import build, write_excel

E18 = 10**18
CHAINS = {
    "base": Chain("base", "Base", 8453, True, ("https://x",), "https://basescan.org"),
    "arc": Chain("arc", "Arc", 5042, True, ("https://y",), "https://explorer.arc.io"),
}
TOKENS = {
    "wARS": Token(
        "wARS",
        "Peso Argentino",
        18,
        None,
        "confirmed",
        "0x0dc4",
        (Deployment("base", "0x0dc4", 10, None), Deployment("arc", "0x0dc4", 500, "0xc0")),
    )
}


def _rec(opening: int, mints: int, burns: int, closing: int) -> dict:
    exp = opening + mints - burns
    return {
        "opening": str(opening),
        "mints": str(mints),
        "burns": str(burns),
        "n_mints": 1 if mints else 0,
        "n_burns": 1 if burns else 0,
        "opening_plus_mints_minus_burns": str(exp),
        "closing": str(closing),
        "difference": str(closing - exp),
        "passed": closing == exp,
    }


def _chain(key: str, opening_basis: str, rec: dict, movements: list[dict]) -> dict:
    blocks = [
        {"cutoff": "2026-06-30", "convention": "ART", "block": 100, "block_timestamp": 1},
        {"cutoff": "2026-09-30", "convention": "ART", "block": 900, "block_timestamp": 2},
    ]
    return {
        "chain": key,
        "previous_cutoff": "2026-06-30",
        "cutoff": "2026-09-30",
        "convention": "ART",
        "cutoff_blocks": blocks,
        "tokens": {
            "wARS": {
                "passed": rec["passed"],
                "source_used": "rpc",
                "sources_complete": ["rpc"],
                "sources_incomplete": [],
                "conventions": {
                    "ART": {
                        "opening_block": 100,
                        "opening_basis": opening_basis,
                        "closing_block": 900,
                        "by_source": {"rpc": rec},
                    }
                },
                "movements": movements,
                "movements_between_conventions": [],
            }
        },
    }


def _mv(chain: str, kind: str, amount: int, category: str) -> dict:
    return {
        "token": "wARS",
        "chain": chain,
        "kind": kind,
        "block": 500,
        "tx_hash": "0x" + "ab" * 32,
        "log_index": 0,
        "amount": str(amount),
        "counterparty": "0x" + "cd" * 20,
        "timestamp": 1_790_000_000,
        "category": category,
        "explorer_url": "https://example/tx",
    }


def test_build_and_excel(tmp_path):
    engine = [
        _chain(
            "base",
            "totalSupply at the previous cutoff block",
            _rec(100 * E18, 5 * E18, 2 * E18, 103 * E18),
            [_mv("base", "mint", 5 * E18, "primary"), _mv("base", "burn", 2 * E18, "unclassified")],
        ),
        _chain(
            "arc",
            "contract created at block 500, after the previous cutoff",
            _rec(0, 7 * E18, 0, 7 * E18),
            [_mv("arc", "mint", 7 * E18, "bridge_in")],
        ),
    ]
    pkg = build(
        engine,
        CHAINS,
        TOKENS,
        {"networks_checked": ["base", "arc"]},
        {"tokens": {"wARS": {"platforms": {"base": "0x0dc4"}}}},
        [{"cutoff": "2026-06-30", "network_keys": ["base"]}],
        {"pairs_matched": 0},
        [{"decision": "d", "label": "documented", "source": "s"}],
        "ART",
    )
    t = pkg["tokens"]["wARS"]
    assert (t["opening"], t["closing"], t["change"]) == ("100", "110", "10")
    assert pkg["all_reconciled"] is True
    assert pkg["new_networks"] == ["arc"]
    assert len(pkg["review"]) == 1 and pkg["review"][0]["kind"] == "burn"
    assert pkg["cross_checks"]["found_not_in_coingecko"] == ["arc"]

    path = tmp_path / "p.xlsx"
    write_excel(pkg, CHAINS, path)
    wb = load_workbook(path)
    assert wb.sheetnames == [
        "Resumen",
        "Por red",
        "Emisiones y quemas",
        "Conciliación",
        "Metodología",
    ]
    assert "Red nueva desde el último cierre: Arc" in [
        c.value for row in wb["Resumen"].iter_rows() for c in row
    ]
